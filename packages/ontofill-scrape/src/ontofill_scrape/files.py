"""Sandbox-client backed, read-only file fetch and offline parsing."""

from __future__ import annotations

import csv
import io
import json
from collections.abc import Mapping
from pathlib import PurePosixPath
from typing import Protocol
from urllib.parse import urljoin, urlsplit

from openpyxl import load_workbook
from pypdf import PdfReader

from .models import FailureKind, FetchedFile, ParsedFile, ParsedRow, ToolFailure


class FetchResponse(Protocol):
    status_code: int
    headers: Mapping[str, str]
    content: bytes


class ReadOnlyFetchClient(Protocol):
    """Implemented by a sandbox browser/proxy; this package makes no network call."""

    def get(self, url: str) -> FetchResponse: ...


def _flatten_json_records(value: object, *, max_rows: int) -> tuple[list[dict[str, object]], bool]:
    """Flatten generic arrays of objects into rows with dotted property paths."""
    truncated = False

    def expand(node: object, prefix: str = "") -> list[dict[str, object]]:
        nonlocal truncated
        if isinstance(node, Mapping):
            records: list[dict[str, object]] = [{}]
            for key, child in node.items():
                child_prefix = f"{prefix}.{key}" if prefix else str(key)
                branches = expand(child, child_prefix)
                merged: list[dict[str, object]] = []
                for left in records:
                    for right in branches:
                        merged.append({**left, **right})
                        if len(merged) > max_rows:
                            truncated = True
                            break
                    if len(merged) > max_rows:
                        break
                records = merged[:max_rows]
            return records
        if isinstance(node, list):
            if node and all(isinstance(item, Mapping) for item in node):
                records = []
                for item in node:
                    records.extend(expand(item, prefix))
                    if len(records) > max_rows:
                        truncated = True
                        return records[:max_rows]
                return records
            return [{prefix: json.dumps(node, ensure_ascii=False, separators=(",", ":"))}]
        return [{prefix: node}] if prefix else [{}]

    rows = expand(value)
    if len(rows) > max_rows:
        truncated = True
        rows = rows[:max_rows]
    return rows, truncated


def _allowed(url: str, allowed_domains: set[str]) -> None:
    parsed = urlsplit(url)
    host = (parsed.hostname or "").lower().rstrip(".")
    if parsed.scheme not in {"http", "https"} or not host:
        raise ToolFailure(FailureKind.BLOCKED, "only public HTTP(S) reads are allowed")
    if parsed.username or parsed.password:
        raise ToolFailure(FailureKind.BLOCKED, "URL credentials are not allowed")
    if not any(host == domain or host.endswith("." + domain) for domain in allowed_domains):
        raise ToolFailure(FailureKind.BLOCKED, f"domain {host} is outside the TDD allowlist")


def file_fetch(
    url: str,
    allowed_domains: set[str] | tuple[str, ...],
    client: ReadOnlyFetchClient,
    *,
    max_bytes: int = 10_000_000,
    max_redirects: int = 3,
) -> FetchedFile:
    domains = {domain.lower().rstrip(".") for domain in allowed_domains}
    if not domains:
        raise ToolFailure(FailureKind.BLOCKED, "TDD allowed_domains is empty")
    current = url
    for _ in range(max_redirects + 1):
        _allowed(current, domains)
        try:
            response = client.get(current)
        except Exception as exc:
            raise ToolFailure(FailureKind.NETWORK, f"fetch failed: {current}") from exc
        if response.status_code in {301, 302, 303, 307, 308}:
            location = response.headers.get("location") or response.headers.get("Location")
            if not location:
                raise ToolFailure(FailureKind.NETWORK, "redirect has no Location")
            current = urljoin(current, location)
            continue
        if response.status_code in {401, 403}:
            raise ToolFailure(FailureKind.BLOCKED, "login or access wall")
        if response.status_code == 429:
            raise ToolFailure(FailureKind.RATE_LIMITED, "source rate limited")
        if response.status_code >= 400:
            raise ToolFailure(FailureKind.NETWORK, f"source returned HTTP {response.status_code}")
        content = response.content
        if len(content) > max_bytes:
            raise ToolFailure(FailureKind.BLOCKED, "file exceeds TDD byte limit")
        content_type = response.headers.get("content-type", "application/octet-stream")
        if b"captcha" in content[:4096].lower() or b"sign in to continue" in content[:4096].lower():
            raise ToolFailure(FailureKind.BLOCKED, "login or captcha wall")
        return FetchedFile(current, content, content_type)
    raise ToolFailure(FailureKind.BLOCKED, "redirect limit exceeded")


def file_parse(
    data: FetchedFile | bytes,
    *,
    format: str | None = None,
    filename: str | None = None,
    max_rows: int = 10_000,
    max_pages: int = 100,
) -> ParsedFile:
    content = data.content if isinstance(data, FetchedFile) else data
    name = filename or (
        PurePosixPath(urlsplit(data.url).path).name if isinstance(data, FetchedFile) else ""
    )
    kind = (format or PurePosixPath(name).suffix.removeprefix(".")).lower()
    if kind == "csv":
        text = content.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:4096])
        except csv.Error:
            dialect = csv.excel
        all_rows = list(csv.reader(io.StringIO(text), dialect))
        rows = tuple(
            ParsedRow(None, index, tuple(row))
            for index, row in enumerate(all_rows[:max_rows], start=1)
            if any(value != "" for value in row)
        )
        return ParsedFile("csv", rows, truncated=len(all_rows) > max_rows)
    if kind in {"xlsx", "xlsm"}:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
        rows: list[ParsedRow] = []
        truncated = False
        total_rows = 0
        try:
            for sheet in workbook.worksheets:
                for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    total_rows += 1
                    if total_rows > max_rows:
                        truncated = True
                        break
                    if any(value is not None and value != "" for value in row):
                        rows.append(ParsedRow(sheet.title, index, tuple(row)))
                if truncated:
                    break
        finally:
            workbook.close()
        return ParsedFile("xlsx", tuple(rows), truncated=truncated)
    if kind == "json":
        try:
            document = json.loads(content)
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ToolFailure(FailureKind.VALIDATION_FAILED, "invalid JSON source") from exc
        records, truncated = _flatten_json_records(document, max_rows=max_rows)
        if not records:
            return ParsedFile("json", ())
        headers = tuple(dict.fromkeys(key for row in records for key in row))
        rows = tuple(
            [ParsedRow(None, 1, headers)]
            + [
                ParsedRow(None, index, tuple(row.get(header) for header in headers))
                for index, row in enumerate(records, start=2)
            ]
        )
        return ParsedFile("json", rows, truncated=truncated)
    if kind == "pdf":
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join((page.extract_text() or "") for page in reader.pages[:max_pages])
        return ParsedFile("pdf", text=text, truncated=len(reader.pages) > max_pages)
    raise ToolFailure(
        FailureKind.VALIDATION_FAILED, f"unsupported file format: {kind or 'unknown'}"
    )
